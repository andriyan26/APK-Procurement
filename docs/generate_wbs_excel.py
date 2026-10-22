"""
Script Generator WBS (Work Breakdown Structure) untuk Soal 1b Take Home Test Cineul Technologies.
Menghasilkan dokumen resmi SOAL_1B_WBS_PROCUREMENT.xlsx dengan styling korporat profesional.
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_wbs():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WBS Procurement System"
    ws.views.sheetView[0].showGridLines = True

    # 1. Judul Header Dokumen
    ws.merge_cells("A1:G1")
    ws["A1"] = "WORK BREAKDOWN STRUCTURE (WBS) - APLIKASI PENGADAAN BARANG & JASA (PROCUREMENT)"
    ws["A1"].font = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:G2")
    ws["A2"] = "Perusahaan: Cineul Technologies | Proyek: Enterprise Procurement Web App (Python Flask & SQLite) | Scope: Request s/d Fisik Gudang"
    ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="475569")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 10  # Blank row

    # 2. Table Header
    headers = [
        "WBS Code",
        "Nama Modul / Fitur / Deliverable",
        "Tingkat (Level)",
        "Deskripsi Fungsional & Spesifikasi Teknis",
        "Peran Penanggung Jawab (Role)",
        "Keluaran (Deliverable / Output)",
        "Status Kesiapan"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws.row_dimensions[4].height = 28
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    # 3. Data WBS
    wbs_data = [
        # Level 1
        ("1.0", "MODUL AUTENTIKASI, KEAMANAN & RBAC", "Level 1 (Modul)", "Sistem manajemen akses terpusat berbasis 5 hak akses peran (RBAC)", "System Admin", "Sistem Login Aman & Sesi Berbasis Cookie", "Selesai (100%)"),
        ("1.1", "Halaman Login & Portal Terpadu", "Level 2 (Fitur)", "Portal login dengan infografis alur pengadaan melingkar & human check anti-bot", "All Roles", "Form Autentikasi dengan CSRF & Session Guard", "Selesai (100%)"),
        ("1.2", "Pemisahan Hak Akses (Role-Based Access Control)", "Level 2 (Fitur)", "Proteksi endpoint @role_required untuk Requester, Approver, Purchasing, Warehouse, Admin", "System Admin", "Middleware & Decorator RBAC", "Selesai (100%)"),
        ("1.3", "Manajemen Akun Pengguna & Profil", "Level 2 (Fitur)", "Pengelolaan kredensial aman (PBKDF2 SHA256), modal edit profil, & ganti password", "All Roles", "Modal Profil Interaktif & Database Users", "Selesai (100%)"),

        # Modul 2
        ("2.0", "MODUL PURCHASE REQUISITION (PR)", "Level 1 (Modul)", "Siklus awal identifikasi dan permohonan pengadaan barang/jasa", "Requester", "Dokumen Elektronik PR & Item Relasional", "Selesai (100%)"),
        ("2.1", "Pengajuan Permohonan PR Baru", "Level 2 (Fitur)", "Formulir interaktif dengan auto-generated nomor PR unik (PR/YYYY/MM/XXXX)", "Requester", "Nomor Registrasi PR & Header Dokumen", "Selesai (100%)"),
        ("2.2", "Multi-Item Grid Dinamis", "Level 2 (Fitur)", "Tambah/hapus baris item kebutuhan tanpa refresh halaman, auto-kalkulasi subtotal & grand total", "Requester", "Tabel Rincian PR (Nama, Spek, Qty, Harga)", "Selesai (100%)"),
        ("2.3", "AI Vision Smart Catalog Scanner (Fitur Nilai Plus)", "Level 2 (Fitur)", "Ekstraksi OCR otomatis dari foto brosur/katalog barang langsung mengisi baris tabel PR", "Requester", "Input Cepat Berbasis AI Vision", "Selesai (100%)"),
        ("2.4", "Monitoring Status & Riwayat PR", "Level 2 (Fitur)", "Tabel filter pills (Semua, Menunggu, Disetujui, Ditolak, PO Terbit) dan pencarian real-time", "Requester", "Daftar Riwayat Pengajuan PR", "Selesai (100%)"),
        ("2.5", "Dokumen Resmi Pengajuan PR", "Level 2 (Fitur)", "Halaman detail dokumen PR siap cetak dengan riwayat status & catatan revisi", "Requester / Approver", "Slip Formal PR Standar A4", "Selesai (100%)"),

        # Modul 3
        ("3.0", "MODUL OTORISASI & APPROVAL WORKFLOW", "Level 1 (Modul)", "Mekanisme tata kelola dan persetujuan anggaran oleh pimpinan divisi", "Approver", "Keputusan Approval & Audit Trail", "Selesai (100%)"),
        ("3.1", "Antrean PR Menunggu Persetujuan", "Level 2 (Fitur)", "Dashboard antrean permohonan prioritas dengan badge urgensi tanggal kebutuhan", "Approver", "Tabel Antrean Otorisasi Pimpinan", "Selesai (100%)"),
        ("3.2", "Peninjauan Rincian Teknis & Anggaran", "Level 2 (Fitur)", "Pemeriksaan detail spesifikasi barang, kewajaran estimasi harga, dan total belanja", "Approver", "Lembar Evaluasi Pengadaan", "Selesai (100%)"),
        ("3.3", "Aksi Keputusan: Approve atau Reject", "Level 2 (Fitur)", "Tombol eksekusi persetujuan atau penolakan wajib menyertakan catatan evaluasi pimpinan", "Approver", "Status Terverifikasi & Log Tanggal Otorisasi", "Selesai (100%)"),
        ("3.4", "Notifikasi Real-Time Feedback", "Level 2 (Fitur)", "Alert status instan untuk memberi kejelasan status dokumen kepada Requester", "Approver / Requester", "Banner Status Approval Interaktif", "Selesai (100%)"),

        # Modul 4
        ("4.0", "MODUL PURCHASE ORDER (PO) & REKANAN VENDOR", "Level 1 (Modul)", "Penerbitan surat pesanan resmi kepada vendor rekanan terpilih", "Purchasing", "Dokumen Kontrak PO & Perjanjian Vendor", "Selesai (100%)"),
        ("4.1", "Antrean PR Siap Terbit PO", "Level 2 (Fitur)", "Daftar permohonan yang telah disetujui (Approved) dan menunggu pemilihan vendor", "Purchasing", "Tabel Seleksi Pengadaan", "Selesai (100%)"),
        ("4.2", "Penerbitan Kontrak PO Baru", "Level 2 (Fitur)", "Formulir kontrak PO lengkap dengan Nomor PO otomatis (PO/YYYY/MM/XXXX)", "Purchasing", "Dokumen PO Terbit", "Selesai (100%)"),
        ("4.3", "AI Vision Quotation Scanner (Fitur Nilai Plus)", "Level 2 (Fitur)", "Ekstraksi OCR surat penawaran resmi vendor otomatis mengisi data vendor & payment terms", "Purchasing", "Input Form Rekanan Cepat via AI", "Selesai (100%)"),
        ("4.4", "Quick Auto-Fill Rekanan Terdaftar", "Level 2 (Fitur)", "Preset satu klik data rekanan terverifikasi (Solusi Mitra, ATK Berkah, Triputra Daya)", "Purchasing", "Pengisian Otomatis Profil Vendor", "Selesai (100%)"),
        ("4.5", "Negosiasi Harga Deal Satuan", "Level 2 (Fitur)", "Penyesuaian harga final kesepakatan dengan live calculator subtotal & grand total PO", "Purchasing", "Tabel Harga Kontrak Pesanan", "Selesai (100%)"),
        ("4.6", "Cetak Dokumen Resmi PO Standar A4", "Level 2 (Fitur)", "Cetak slip PO formal lengkap dengan Kop Surat Perusahaan, Cap Stempel Digital, & 3 Tanda Tangan", "Purchasing / Vendor", "Dokumen Legal PO Siap Kirim (PDF Print)", "Selesai (100%)"),

        # Modul 5
        ("5.0", "MODUL GOODS RECEIPT (GR) & FISIK GUDANG", "Level 1 (Modul)", "Verifikasi penerimaan fisik barang dan surat jalan di gudang logistik", "Warehouse", "Surat Bukti Penerimaan Barang (GR Slip)", "Selesai (100%)"),
        ("5.1", "Antrean PO Menunggu Kedatangan Fisik", "Level 2 (Fitur)", "Monitoring Purchase Order berstatus aktif (Issued) yang belum diterima di gudang", "Warehouse", "Daftar Kedatangan Kargo Logistik", "Selesai (100%)"),
        ("5.2", "AI Computer Vision: Smart Scanner Surat Jalan (Fitur Nilai Plus)", "Level 2 (Fitur)", "Pemindaian dokumen Delivery Order (DO) fisik kurir untuk ekstraksi nomor DO & kondisi segel", "Warehouse", "Input Data Surat Jalan Otomatis", "Selesai (100%)"),
        ("5.3", "Pemeriksaan Kuantitas & Kondisi Fisik", "Level 2 (Fitur)", "Verifikasi kuantitas riil per item serta tagging kondisi: Baik, Rusak, atau Salah Tipe", "Warehouse", "Berita Acara Pemeriksaan Fisik", "Selesai (100%)"),
        ("5.4", "Kalkulasi Otomatis Status Penerimaan", "Level 2 (Fitur)", "Sistem otomatis menetapkan status Received (Lengkap) atau Partially Received (Sebagian)", "Warehouse", "Update Status PO & GR Real-Time", "Selesai (100%)"),
        ("5.5", "Cetak Lembar Tanda Terima Goods Receipt", "Level 2 (Fitur)", "Pencetakan slip serah terima barang gudang resmi ber-Kop Surat dan tanda tangan penerima", "Warehouse", "Dokumen Fisik GR Berita Acara", "Selesai (100%)"),

        # Modul 6
        ("6.0", "MODUL ANALITIK, DASHBOARD & PELAPORAN", "Level 1 (Modul)", "Pusat informasi KPI dan rekapitulasi performa pengadaan organisasi", "All Roles", "Laporan Manajemen & Grafik Tren", "Selesai (100%)"),
        ("6.1", "Role-Tailored Dashboard Eksekutif", "Level 2 (Fitur)", "Tampilan metrik statistik disesuaikan khusus untuk tiap peran yang sedang login", "All Roles", "Dasbor Interaktif Khusus Role", "Selesai (100%)"),
        ("6.2", "Visualisasi Grafik Real-Time (Chart.js)", "Level 2 (Fitur)", "Grafik distribusi status pengadaan dan tren belanja anggaran per divisi", "Approver / Admin", "Widget Visual Tren Pengadaan", "Selesai (100%)"),
        ("6.3", "Audit Trail & Rekap Riwayat Transaksi", "Level 2 (Fitur)", "Pelacakan alur dokumen hulu-hilir (PR &rarr; Approval &rarr; PO &rarr; GR) secara transparan", "Admin / Auditor", "Log Aktivitas Lengkap Sistem", "Selesai (100%)"),

        # Modul 7
        ("7.0", "FITUR OPSIONAL NILAI PLUS: AI VISION SUITE", "Level 1 (Modul)", "Implementasi Computer Vision & AI Model terpadu untuk efisiensi form input", "Cross-Role", "Suite Asisten AI Pengadaan Cerdas", "Selesai (100%)"),
        ("7.1", "AI Smart Catalog Extractor (Requester)", "Level 2 (Fitur)", "Ekstraksi multi-item otomatis dari foto katalog/brosur ke form PR", "Requester", "Pengisian Otomatis Item PR via AI", "Selesai (100%)"),
        ("7.2", "AI Quotation OCR Extractor (Purchasing)", "Level 2 (Fitur)", "Ekstraksi otomatis lembar surat penawaran vendor ke form PO", "Purchasing", "Pengisian Otomatis Vendor PO via AI", "Selesai (100%)"),
        ("7.3", "AI Delivery Order Scanner (Warehouse)", "Level 2 (Fitur)", "Ekstraksi nomor DO & deteksi kondisi segel fisik paket pengiriman", "Warehouse", "Pengisian Otomatis Form GR via AI", "Selesai (100%)")
    ]

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    module_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    current_row = 5
    for row_data in wbs_data:
        wbs_code, name, level, desc, role, deliverable, status = row_data
        ws.row_dimensions[current_row].height = 24

        is_level1 = "Level 1" in level

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.border = data_border

            if is_level1:
                cell.fill = module_fill
                cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="1E3A8A")
            else:
                if current_row % 2 == 0:
                    cell.fill = zebra_fill
                cell.font = Font(name="Segoe UI", size=9, bold=(col_idx == 2 and is_level1), color="0F172A")

            # Alignment
            if col_idx in [1, 3, 5, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Set Column Widths
    col_widths = {
        1: 14,  # WBS Code
        2: 38,  # Nama Modul
        3: 18,  # Tingkat
        4: 55,  # Deskripsi
        5: 22,  # Peran
        6: 35,  # Output
        7: 18   # Status
    }

    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    os.makedirs("docs", exist_ok=True)
    output_path = os.path.join("docs", "SOAL_1B_WBS_PROCUREMENT.xlsx")
    wb.save(output_path)
    print(f"Berhasil membuat berkas Excel: {output_path}")

if __name__ == "__main__":
    generate_wbs()
